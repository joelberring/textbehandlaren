import sys
from docxtpl import DocxTemplate
from docx import Document

def inspect_docx(path):
    print(f"Inspecting: {path}")
    try:
        doc = DocxTemplate(path)
        # Find placeholders
        placeholders = doc.get_undeclared_template_variables()
        print(f"Placeholders found: {placeholders}")
        
        # List styles
        base_doc = Document(path)
        print("\nAvailable Paragraph Styles:")
        for style in base_doc.styles:
            if style.type == 1: # Paragraph style
                print(f"- {style.name}")
                
        # Sample content
        print("\nFirst 10 Paragraphs (Text + Style):")
        for i, p in enumerate(base_doc.paragraphs[:10]):
            print(f"{i}: [{p.style.name}] {p.text[:100]}...")
            
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        inspect_docx(sys.argv[1])
    else:
        print("Usage: python3 inspect_docx.py <path>")
